#!/usr/bin/env python3
"""
xlsx_to_knowledge_md.py — turn a workbook into corpus markdown, by transposing.

The third converter, for the documents that are spreadsheets. It exists because
the obvious approach fails badly on this corpus.

## Why not just read the text

A Drive text export of the customer journey map arrives as 244,000 characters
on a single line, with cell values separated by commas — and the cell values
themselves contain commas ("Ask 3D Personnel if they can supply and we process
payment after 30 working days"). Splitting that on commas shatters sentences
and, worse, shifts every subsequent column by one. In a journey map the column
IS the meaning: an action filed under the wrong stage is a confident, specific,
wrong answer about what a customer does. So this reads the real workbook
through openpyxl, where a cell is a cell.

## Why transpose

A journey map is stored with journey stages across the top and dimensions down
the side. Read row-wise, a chunk is "every action, at every stage" — useless.
Read column-wise, a chunk is one stage of one journey with its actions,
questions, pain points and opportunities together, which is what somebody is
actually asking about.

## Shapes

Three, detected per sheet, because these five sheets use three layouts:

  journey-col-label   Band labels merged down column A; stages in row 1,
                      steps in row 2. (OnBoarding_Master, OFFBoarding_Churn)
  journey-row-label   Band labels on their own row; stages and steps in
                      alternating columns. (After_Sales_Service)
  table               An ordinary header row with data rows beneath.
                      (Responsibility Matrix v1.0, v2.0)

Detection is reported per sheet and can be overridden with --shape, because a
wrong guess here is silent: it produces plausible markdown from misread cells.
Run with --dry-run first and read the preview.

    pip install openpyxl
    python scripts/xlsx_to_knowledge_md.py book.xlsx --name customer_journey_maps \\
        --title "Customer Journey Maps" --dry-run
    python scripts/xlsx_to_knowledge_md.py book.xlsx --name customer_journey_maps \\
        --title "Customer Journey Maps" -o database/knowledge
"""

import argparse
import os
import re
import unicodedata

try:
    import openpyxl
except ImportError:
    raise SystemExit("openpyxl is required:  pip install openpyxl")

MAX_SCAN_ROWS = 400        # these sheets declare 1000+ rows, nearly all empty
MAX_SCAN_COLS = 60


def text(value):
    if value is None:
        return ''
    out = unicodedata.normalize('NFKC', str(value)).replace(' ', ' ')
    out = out.replace('\r', ' ').replace('\n', ' / ')
    return re.sub(r'\s+', ' ', out).strip()


def grid(worksheet):
    """A dense [row][col] grid of strings, with merged cells filled in.

    openpyxl gives a merged range's value only in its top-left cell and None
    everywhere else. For a header merged across six columns that means five
    blank stage names, so the fill is not cosmetic.
    """
    rows = min(worksheet.max_row or 0, MAX_SCAN_ROWS)
    cols = min(worksheet.max_column or 0, MAX_SCAN_COLS)
    table = [['' for _ in range(cols + 1)] for _ in range(rows + 1)]
    for r in range(1, rows + 1):
        for c in range(1, cols + 1):
            table[r][c] = text(worksheet.cell(row=r, column=c).value)

    for merged in worksheet.merged_cells.ranges:
        if merged.min_row > rows or merged.min_col > cols:
            continue
        value = table[merged.min_row][merged.min_col]
        if not value:
            continue
        for r in range(merged.min_row, min(merged.max_row, rows) + 1):
            for c in range(merged.min_col, min(merged.max_col, cols) + 1):
                table[r][c] = value
    return table, rows, cols


def title_case(value):
    """Title case that leaves the letter after an apostrophe alone.

    str.title() renders "QUESTIONS IN BUYER'S MIND" as "Buyer'S Mind", which
    ends up in the chunk a person reads.
    """
    return re.sub(r"(?<![A-Za-z'])([a-z])", lambda m: m.group(1).upper(),
                  value.lower())


def is_band(value):
    """A dimension label: ALL CAPS, a few words, not a sentence."""
    if not value or len(value) > 70:
        return False
    letters = [ch for ch in value if ch.isalpha()]
    if len(letters) < 3:
        return False
    return sum(ch.isupper() for ch in letters) / len(letters) > 0.85


def detect(table, rows, cols):
    band_in_col_a = sum(1 for r in range(1, rows + 1) if is_band(table[r][1]))
    band_on_own_row = sum(
        1 for r in range(1, rows + 1)
        if not table[r][1] and is_band(table[r][2])
        and not any(table[r][c] for c in range(3, cols + 1)))
    if band_in_col_a >= 2:
        return 'journey-col-label'
    if band_on_own_row >= 2:
        return 'journey-row-label'
    return 'table'


def header_row(table, rows, cols):
    """The first row that looks like column headings: several short cells."""
    for r in range(1, min(rows, 12) + 1):
        filled = [table[r][c] for c in range(1, cols + 1) if table[r][c]]
        if len(filled) >= 3 and all(len(v) <= 60 for v in filled):
            return r
    return 1


def journey_col_label(table, rows, cols):
    """Stages in row 1, steps in row 2, band labels merged down column A."""
    steps = []
    for c in range(2, cols + 1):
        stage, step = table[1][c], table[2][c]
        if not step:
            continue
        bands, current = [], None
        for r in range(3, rows + 1):
            label, value = table[r][1], table[r][c]
            if label and (not current or current[0] != label):
                current = (label, [])
                bands.append(current)
            if value and current:
                current[1].append(value)
        steps.append((stage, step, [(l, v) for l, v in bands if v]))
    return steps


def journey_row_label(table, rows, cols):
    """Band labels on their own row; stages and steps in alternating columns."""
    stage_row = next((r for r in range(1, min(rows, 8) + 1)
                      if sum(1 for c in range(1, cols + 1) if table[r][c]) >= 2), 1)
    step_row = next((r for r in range(stage_row + 1, min(rows, 10) + 1)
                     if sum(1 for c in range(1, cols + 1) if table[r][c]) >= 2),
                    stage_row + 1)

    band_rows = [r for r in range(step_row + 1, rows + 1)
                 if not table[r][1] and is_band(table[r][2])
                 and not any(table[r][c] for c in range(3, cols + 1))]

    steps = []
    for c in range(1, cols + 1):
        step = table[step_row][c]
        if not step:
            continue
        stage = table[stage_row][c]
        bands = []
        for index, start in enumerate(band_rows):
            stop = band_rows[index + 1] if index + 1 < len(band_rows) else rows + 1
            values = [table[r][c] for r in range(start + 1, stop) if table[r][c]]
            if values:
                bands.append((table[start][2], values))
        if bands:
            steps.append((stage, step, bands))
    return steps


def as_table(table, rows, cols):
    head = header_row(table, rows, cols)
    headings = [table[head][c] for c in range(1, cols + 1)]
    records = []
    for r in range(head + 1, rows + 1):
        pairs = [(headings[c - 1] or f'Column {c}', table[r][c])
                 for c in range(1, cols + 1) if table[r][c]]
        if pairs:
            records.append(pairs)
    return records


def render(worksheet, shape):
    table, rows, cols = grid(worksheet)
    lines = [f'## {worksheet.title.strip()}', '']

    if shape == 'table':
        records = as_table(table, rows, cols)
        for record in records:
            lines.append(f'### {record[0][1][:90]}')
            for label, value in record:
                lines.append(f'- **{label}**: {value}')
            lines.append('')
        return lines, len(records)

    steps = (journey_col_label(table, rows, cols) if shape == 'journey-col-label'
             else journey_row_label(table, rows, cols))
    stage_seen = None
    for stage, step, bands in steps:
        if stage and stage != stage_seen:
            lines += [f'### {stage}', '']
            stage_seen = stage
        lines += [f'#### {step}', '']
        for label, values in bands:
            lines.append(f'**{title_case(label)}**')
            lines += [f'- {v}' for v in values]
            lines.append('')
    return lines, len(steps)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('source', help='the .xlsx workbook')
    parser.add_argument('-o', '--out', default='database/knowledge')
    parser.add_argument('--name', required=True, help='output filename stem')
    parser.add_argument('--title', help='H1 for the document')
    parser.add_argument('--shape', choices=['journey-col-label',
                                            'journey-row-label', 'table'],
                        help='override detection for every sheet')
    parser.add_argument('--sheet', action='append',
                        help='only this sheet (repeatable)')
    parser.add_argument('--dry-run', action='store_true',
                        help='report what was detected and preview, write nothing')
    args = parser.parse_args()

    workbook = openpyxl.load_workbook(args.source, data_only=True)
    out = [f"# {args.title or args.name}", '']
    total = 0

    for worksheet in workbook.worksheets:
        if args.sheet and worksheet.title.strip() not in [s.strip() for s in args.sheet]:
            continue
        table, rows, cols = grid(worksheet)
        shape = args.shape or detect(table, rows, cols)
        lines, count = render(worksheet, shape)
        total += count
        print(f"  {worksheet.title.strip()!r}: {shape}, {count} "
              f"{'record' if shape == 'table' else 'journey step'}(s)")
        if count == 0:
            print(f"     CHECK: nothing extracted. Wrong shape? Override with "
                  f"--shape and --sheet.")
        out += lines

    body = re.sub(r'\n{3,}', '\n\n', '\n'.join(out)).strip() + '\n'

    if args.dry_run:
        print(f"\n--- preview, first 1500 chars of {len(body):,} ---\n")
        print(body[:1500])
        print("\nDry run — nothing written.")
        return

    os.makedirs(args.out, exist_ok=True)
    target = os.path.join(args.out, f'{args.name}.md')
    with open(target, 'w', encoding='utf-8') as handle:
        handle.write(body)
    print(f"\nwrote {target}\n  {len(body):,} chars, {total} sections")
    print("  Add it to database/knowledge/manifest.json, then:")
    print(f"  python scripts/ingest_waswa_knowledge.py --only {args.name}.md --commit")


if __name__ == '__main__':
    main()
